from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

# Column aliases to resolve inconsistent headings
COLUMN_MAPPING: Dict[str, List[str]] = {
    "date": [
        "date",
        "transaction_date",
        "date_of_transaction",
        "voucher_date",
        "dt",
        "trans_date",
    ],
    "amount": [
        "amount",
        "value",
        "transaction_amount",
        "amt",
        "amount_dr",
        "amount_cr",
        "debit",
        "credit",
        "dr",
        "cr",
    ],
    "ledger": ["ledger", "account", "account_name", "ledger_name", "particulars"],
    "customer": ["customer", "vendor", "party", "customer_name", "party_name"],
    "voucher_type": ["voucher_type", "type", "transaction_type", "voucher"],
    "description": ["description", "memo", "notes", "narration"],
    "reference": ["reference", "ref", "cheque_number", "cheque", "ref_no"],
    "quantity": ["quantity", "qty", "units"],
    "item": ["item", "product", "item_name", "stock_item"],
    "cost_centre": ["cost_centre", "cost_center", "department", "project"],
    "debit": ["debit", "dr", "amount_dr"],
    "credit": ["credit", "cr", "amount_cr"],
}

MONTH_MAP = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}


class ExcelLoadError(RuntimeError):
    pass


def _ensure_directory_exists(directory: Path) -> None:
    if not directory.exists():
        raise ExcelLoadError(f"Directory not found or unreachable: {directory}")
    if not directory.is_dir():
        raise ExcelLoadError(f"Configured path is not a directory: {directory}")


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Normalize column names to snake_case and trim whitespace."""
    df = df.copy()
    df.columns = [re.sub(r"[\s\-]+", "_", col.lower().strip()) for col in df.columns]
    return df


def _find_column(df: pd.DataFrame, key: str) -> Optional[str]:
    """Find a column by alias list, returning the actual column name."""
    aliases = COLUMN_MAPPING.get(key, [])
    normalized = {col.lower(): col for col in df.columns}
    for alias in aliases:
        if alias in normalized:
            return normalized[alias]
    return None


def _resolve_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Rename columns to canonical names; raise if required are missing."""
    df = _normalize_columns(df)
    renamed: Dict[str, str] = {}
    required = ["date", "amount", "ledger"]
    optional = [
        "customer",
        "voucher_type",
        "description",
        "reference",
        "quantity",
        "item",
        "cost_centre",
        "debit",
        "credit",
    ]

    for key in required + optional:
        found = _find_column(df, key)
        if found:
            renamed[found] = key

    missing_required = [key for key in required if key not in renamed.values()]
    if missing_required:
        # Try to derive amount from debit/credit if missing amount
        if "amount" in missing_required:
            has_debit = any(alias in df.columns for alias in COLUMN_MAPPING["debit"])
            has_credit = any(alias in df.columns for alias in COLUMN_MAPPING["credit"])
            if has_debit and has_credit:
                missing_required.remove("amount")
            else:
                # Heuristic: pick first column with amount-ish name
                amt_like = next((c for c in df.columns if any(tok in c.lower() for tok in ["amt", "amount", "value", "net"])), None)
                if amt_like:
                    renamed[amt_like] = "amount"
                    missing_required = [m for m in missing_required if m != "amount"]
        if "date" in missing_required:
            # Heuristic: pick first column containing 'date'
            auto_date = next((c for c in df.columns if "date" in c.lower()), None)
            if auto_date:
                renamed[auto_date] = "date"
                missing_required = [m for m in missing_required if m != "date"]
        if missing_required:
            raise ExcelLoadError(
                f"Missing required columns: {', '.join(missing_required)}. Available: {', '.join(df.columns)}"
            )

    df = df.rename(columns=renamed)

    for key in optional:
        if key not in df.columns:
            df[key] = ""

    return df


def _maybe_melt_monthly(df: pd.DataFrame) -> Tuple[pd.DataFrame, bool]:
    """If monthly columns like apr_25 exist, melt them into long format with date/amount."""
    month_pattern = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)_\d{2}$")
    month_cols = [c for c in df.columns if month_pattern.match(c.lower())]
    if not month_cols:
        return df, False

    id_vars = [c for c in df.columns if c not in month_cols and c != "total"]

    def _to_date(key: str) -> pd.Timestamp:
        lower = key.lower()
        mon, yy = lower.split("_")
        year = 2000 + int(yy)
        month = MONTH_MAP.get(mon)
        return pd.Timestamp(year=year, month=month, day=1)

    melted = df.melt(id_vars=id_vars, value_vars=month_cols, var_name="month_key", value_name="amount")
    melted["date"] = melted["month_key"].apply(_to_date)
    melted = melted.dropna(subset=["amount"])
    melted = melted.drop(columns=["month_key"], errors="ignore")
    return melted, True


def _coerce_types(df: pd.DataFrame) -> pd.DataFrame:
    """Coerce date and amount columns; drop rows that cannot convert."""
    df = df.copy()
    # If debit/credit present but amount missing, derive amount = credit - debit
    if "amount" not in df.columns and "debit" in df.columns and "credit" in df.columns:
        df["amount"] = pd.to_numeric(df["credit"], errors="coerce") - pd.to_numeric(df["debit"], errors="coerce")

    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    df["amount"] = pd.to_numeric(df["amount"], errors="coerce")
    df = df.dropna(subset=["date", "amount", "ledger"])
    df["ledger"] = df["ledger"].astype(str).str.strip()
    df["customer"] = df["customer"].astype(str).str.strip()
    return df


def _pick_sheet(workbook_path: Path, default_sheet: str, fallbacks: Iterable[str]) -> str:
    wb = load_workbook(workbook_path, read_only=True)
    if default_sheet in wb.sheetnames:
        return default_sheet
    for name in fallbacks:
        if name in wb.sheetnames:
            return name
    # Heuristic: prefer first sheet containing 'trans' or 'report'
    for name in wb.sheetnames:
        lower = name.lower()
        if "trans" in lower or "report" in lower:
            return name
    return wb.sheetnames[0]


def discover_latest_file(directory: Path, pattern: str) -> Path:
    """Find the latest file by modification time matching pattern."""
    _ensure_directory_exists(directory)
    candidates = sorted(directory.glob(pattern))
    if not candidates:
        raise ExcelLoadError(f"No files matching pattern '{pattern}' in {directory}")
    return max(candidates, key=lambda p: p.stat().st_mtime)


def latest_cache_key(excel_config: Dict[str, object]) -> Tuple[str, float]:
    """Return (path, mtime) for the newest file to drive cache invalidation."""
    smb_share = Path(str(excel_config.get("smb_share", "")))
    pattern = str(excel_config.get("file_pattern", "*.xlsx"))
    latest_path = discover_latest_file(smb_share, pattern)
    return str(latest_path), latest_path.stat().st_mtime


def load_excel_at_path(excel_path: Path, excel_config: Dict[str, object]) -> Tuple[pd.DataFrame, Dict[str, object]]:
    """Load and clean Excel from an explicit path (used with cache key)."""
    default_sheet = str(excel_config.get("default_sheet", "Transactions"))
    fallbacks = excel_config.get("fallback_sheets", []) or []

    if not excel_path.exists():
        raise ExcelLoadError(f"Excel file not found: {excel_path}")

    sheet_name = _pick_sheet(excel_path, default_sheet, fallbacks)

    # Detect header row
    header_idx = int(excel_config.get("header_row", 0))
    
    # Try to auto-detect if default 0 fails (common for Tally exports)
    if header_idx == 0:
        preview = pd.read_excel(excel_path, sheet_name=sheet_name, header=None, nrows=10, engine="openpyxl")
        for i, row in preview.iterrows():
            row_vals = [str(x).lower() for x in row.tolist()]
            if "date" in row_vals and "ledger" in row_vals:
                header_idx = i
                break

    df = pd.read_excel(excel_path, sheet_name=sheet_name, header=header_idx, engine="openpyxl")
    df = _normalize_columns(df)
    df, melted = _maybe_melt_monthly(df)
    df = _resolve_columns(df)
    df = _coerce_types(df)
    df = df.sort_values("date")

    metadata = {
        "file_path": str(excel_path),
        "sheet": sheet_name,
        "rows": int(len(df)),
        "columns": list(df.columns),
        "modified_at": datetime.fromtimestamp(excel_path.stat().st_mtime),
    }
    return df, metadata


def load_latest_excel(excel_config: Dict[str, object]) -> Tuple[pd.DataFrame, Dict[str, object]]:
    """Backward-compatible helper: discover then load."""
    latest_path = discover_latest_file(Path(str(excel_config.get("smb_share", ""))), str(excel_config.get("file_pattern", "*.xlsx")))
    return load_excel_at_path(latest_path, excel_config)
